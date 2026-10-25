"""Create the slide deck and speaker notes from docs/results.json."""

import json
import re
from pathlib import Path

from PIL import Image
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

from .config import DOCS, ROOT

RESULTS = DOCS / "results.json"
SLIDES_DIR = ROOT / "slides"


def n(value, digits=2):
    return f"{value:,.{digits}f}"


def pct(value, digits=1):
    return f"{100 * value:.{digits}f}%"


def p(value):
    return f"{value:.2e}" if value < 0.001 else f"{value:.3f}"


def figure(name):
    return ROOT / "reports" / "figures" / name


def content_figures(content):
    if "figures" in content:
        return content["figures"]
    return [content["figure"]] if content.get("figure") else []


def build_deck(data):
    ds, target, age = data["dataset"], data["target"], data["age"]
    integrity = data["integrity"]
    views, h1, h4 = data["views"], data["h1"], data["h4"]
    spec = data["external_context"]["uspstf"]
    uspstf_ref = data["external_context"]["references"][0]
    deliverable = data["deliverable_spec"]

    slides = [
        {
            "title": "The file’s structure dominates its label associations",
            "body": [
                f"{ds['n_raw']:,} rows, {ds['n_columns']} columns; source is undocumented.",
                f"After removing the ordered block, smoking OR is {n(h1['D_dedup_nolead']['odds_ratio'], 2)}.",
                "Interpretation: integrity audit and descriptive analysis only.",
            ],
            "notes": "Open with the single takeaway. This is a review of one supplied file, not a population study. State that no model was fitted and that the checksum mismatch is explicitly handled.",
        },
        {
            "title": "A smoking checkbox cannot reproduce screening history",
            "body": [
                f"USPSTF {spec['year']} [{uspstf_ref['number']}] uses ages {spec['age_min']}–{spec['age_max']}.",
                f"It also requires at least {spec['pack_years_min']} pack-years and a quit-time criterion of {spec['quit_years_max']} years.",
                "This file has no dose, duration, or quit-time fields; no eligibility inference is possible.",
            ],
            "notes": "Explain the reference criteria and immediately contrast them with the available fields. This is context for limitations, not an assessment of any person.",
        },
        {
            "title": "The objective is to test whether findings survive integrity checks",
            "body": [
                "RQ1: characterize ordering, duplicates, and conflicting records.",
                "RQ2–RQ3: quantify associations and their sensitivity to cleaning.",
                "RQ4–RQ5: describe representativeness limits and justify candidate features.",
                "Success means reproducible checks, transparent uncertainty, and traceable claims—not model accuracy.",
            ],
            "notes": "State that the project is descriptive and integrity-focused. The five research questions are fully listed in the report; summarize the paired questions here rather than reading them verbatim.",
        },
        {
            "title": "The checksum does not match the project contract",
            "body": [
                f"Observed: {ds['n_raw']:,} × {ds['n_columns']}; {ds['null_count']} null cells.",
                f"Expected SHA-256 prefix: {ds['expected_sha256_prefix']}.",
                f"Observed prefix: {ds['sha256_prefix']}.",
                f"Recorded-label code 1 share: {pct(target['positive_rate'])}; not population prevalence.",
            ],
            "notes": "Describe the mismatch plainly. The user approved continuing with this current snapshot under an explicit override; the expected checksum was not changed.",
        },
        {
            "title": "Several meanings needed for interpretation are unknown",
            "body": [
                "Dataset source, sampling frame, collection process, and licence: unknown.",
                "Meaning and generation of the recorded label: unknown.",
                "GENDER code mapping: unknown.",
                "The project README only supplied a title; redistribution needs a licence check.",
            ],
            "notes": "Do not speculate about where the file came from or what the label means. Point to the data card and open questions.",
        },
        {
            "title": "The leading rows form an all-negative block",
            "body": [
                f"First {integrity['leading_block_n']:,} rows are coded negative.",
                f"Smoking YES share: {pct(integrity['source_shift_smoking_lead_yes'])} in the block vs {pct(integrity['source_shift_smoking_later_yes'])} in later negatives.",
                "The plot shows row order, not a documented timeline.",
            ],
            "figures": [figure("f1_order.png"), figure("v6_yes_share_by_block.png")],
            "figure_refs": ["Figure 1: row order", "Figure 2: YES-share profiles"],
            "notes": "Figure 1: describe row index, rolling positive-label share, and the boundary. Figure 2: compare YES shares across leading negatives, later negatives, and later positives. Both are descriptive; neither establishes provenance.",
        },
        {
            "title": "Duplicate patterns and ID order are not harmless",
            "body": [
                f"{integrity['exact_duplicates']:,} exact duplicate rows; {integrity['conflicting_patterns']:,} conflicting feature patterns.",
                f"Maximum pattern multiplicity: {integrity['max_pattern_multiplicity']}.",
                f"ID-only AUC: {n(integrity['patient_id_auc'], 2)} — an audit statistic, not model performance.",
                "The source and record-generation process remain unknown.",
            ],
            "figure": figure("v4_duplicate_multiplicity.png"),
            "figure_refs": ["Duplicate multiplicity and conflicting patterns"],
            "notes": "Describe the duplicate-multiplicity distribution and conflicting-pattern counts. Explain why duplicate policies are sensitivity analyses. The ID-only AUC is an integrity statistic; patient_id remains excluded from feature sets.",
        },
        {
            "title": "Four cleaning views expose policy sensitivity",
            "body": [
                f"A raw: {views['A_raw']['n']:,}; B exact-deduplicated: {views['B_dedup']['n']:,}.",
                f"C no conflicts: {views['C_dedup_noconflict']['n']:,}; D no leading block: {views['D_dedup_nolead']['n']:,}.",
                f"Recorded-positive share ranges from {pct(views['A_raw']['positive_rate'])} to {pct(views['D_dedup_nolead']['positive_rate'])}.",
                "View D filters B by the original raw row index.",
            ],
            "figure": figure("fig5_views.png"),
            "figure_refs": ["Figure 5: four cleaning views"],
            "notes": "Figure 5 shows retained row counts and recorded-label-1 rates for views A through D. Explain that these are sensitivity policies, not a known ground-truth repair.",
        },
        {
            "title": "After deduplication, pairwise feature relations are weak",
            "body": [
                f"Age: {age['min']}–{age['max']}, mean {n(age['mean'], 1)}.",
                f"Maximum absolute pairwise correlation: {n(integrity['independence']['max_abs_pairwise_corr'], 3)}.",
                f"First PCA component: {pct(integrity['independence']['pca_first_component'])}; flat reference: {pct(integrity['independence']['pca_flat_reference'])}.",
                "These are descriptive diagnostics, not proof of independence.",
            ],
            "figures": [figure("fig3_correlation.png"), figure("v7_age_profiles.png")],
            "figure_refs": ["Figure 3: feature correlations", "Figure 4: AGE profiles"],
            "notes": "Figure 3: describe the feature-pair axes, correlation colour, and clipped range. Figure 4: compare AGE across the leading block, later negatives, and later positives, and show the label-wise and AGE_BAND summaries. These are descriptive views.",
        },
        {
            "title": "Five pre-specified tests separate effect from significance",
            "body": [
                "H1 smoking: odds ratio, confidence interval, chi-square, permutation check.",
                "H2 age: Mann–Whitney and CLES; H3 symptom count: Spearman rho.",
                "H4 feature family: Benjamini–Hochberg correction; H5: source-shift audit.",
                f"View D H4 minimum raw p = {p(h4['D_dedup_nolead']['min_p'])}; minimum q = {n(h4['D_dedup_nolead']['min_q'], 3)}.",
            ],
            "notes": "Describe the pre-specified tests and distinguish p-values, effect sizes, and adjusted q-values. The analysis does not test causal hypotheses.",
        },
        {
            "title": "The smoking association moves toward the null after cleaning",
            "body": [
                f"View A OR {n(h1['A_raw']['odds_ratio'], 2)} ({n(h1['A_raw']['ci_low'], 2)}–{n(h1['A_raw']['ci_high'], 2)}).",
                f"View D OR {n(h1['D_dedup_nolead']['odds_ratio'], 2)} ({n(h1['D_dedup_nolead']['ci_low'], 2)}–{n(h1['D_dedup_nolead']['ci_high'], 2)}).",
                f"View D H4 minimum q = {n(h4['D_dedup_nolead']['min_q'], 2)}.",
                "The file does not establish a population or causal conclusion.",
            ],
            "figures": [figure("fig7_smoking_forest.png"), figure("fig6_cramers_v.png")],
            "figure_refs": ["Figure 7: smoking odds ratio by view", "Figure 6: Cramer's V by feature"],
            "notes": "Start with Figure 7: describe view, smoking odds ratio, 95% intervals, and the OR=1 reference. Then use Figure 6 to show the feature-level Cramer's V comparison from A to D. These are associations in this file, not population conclusions.",
        },
        {
            "title": "Feature construction is transparent, not evidence of prediction",
            "body": [
                f"F0: {data['features']['F0_raw15']} raw features; F1: {data['features']['F1_raw_plus_engineered']} with selected descriptors.",
                f"F2: {data['features']['F2_prior_driven']} domain-prior candidates.",
                "Age scaling must be fitted inside training folds in any future model.",
                "The smoking-by-age feature is a proxy, not screening eligibility.",
            ],
            "notes": "Summarize the three sets and the rationale for retaining F2 as a candidate. State that univariate screening did not select features and no model was fitted.",
        },
        {
            "title": "The file supports an integrity finding, not a clinical claim",
            "body": [
                f"After removing the leading block, smoking OR is {n(h1['D_dedup_nolead']['odds_ratio'], 2)}.",
                f"After BH correction in D, the minimum q is {n(h4['D_dedup_nolead']['min_q'], 2)}.",
                "Provenance, label meaning, representativeness, and external validity are unresolved.",
                "All conclusions are limited to this supplied file.",
            ],
            "notes": "Answer the central question directly. Explain that small effects after structural cleaning are about this file only and do not negate external evidence.",
        },
        {
            "title": "Next steps are planned, not completed",
            "body": [
                "Verify source, licence, checksum, label definition, and sampling frame.",
                "If authorized, plan grouped/source-aware validation on independent data.",
                "Select metrics and calibration checks with the intended use in mind.",
                "Clean-clone testing and timed presentation rehearsal remain to be done.",
            ],
            "notes": "Use future tense for all uncompleted work. Do not claim that external validation, model training, USB backup, or rehearsal has occurred.",
        },
    ]

    backup = [
        {
            "title": "Backup: sensitivity estimates by view",
            "body": [
                f"{label}: OR {n(h1[label]['odds_ratio'], 2)}, 95% CI {n(h1[label]['ci_low'], 2)}–{n(h1[label]['ci_high'], 2)}, p {p(h1[label]['p'])}."
                for label in ["A_raw", "B_dedup", "C_dedup_noconflict", "D_dedup_nolead"]
            ],
            "notes": "Backup sensitivity values. Keep interpretation descriptive and point to the integrity audit.",
        },
        {
            "title": "Backup: feature-family BH results",
            "body": [
                f"{label}: minimum p {p(h4[label]['min_p'])}, minimum q {n(h4[label]['min_q'], 3)}, maximum V {n(h4[label]['max_v'], 3)}."
                for label in ["A_raw", "B_dedup", "C_dedup_noconflict", "D_dedup_nolead"]
            ],
            "notes": "Backup multiple-testing family summary. The H4 tests are descriptive, not evidence of causation.",
        },
        {
            "title": "Backup: detectable risk-difference approximation",
            "body": [
                f"View D group sizes: {data['power']['view_D_smokers_n']:,} smoking-code YES and {data['power']['view_D_nonsmokers_n']:,} NO.",
                f"Alpha {n(data['power']['alpha'], 2)}; target power {n(data['power']['power'], 1)}.",
                f"Approximate minimum detectable risk difference: {n(data['power']['minimum_detectable_risk_difference'], 4)}.",
                "Planning calculation only; its assumptions are not established by the file.",
            ],
            "notes": "Explain that this is a planning calculation under assumptions, not evidence that a smaller association is absent.",
        },
        {
            "title": "Backup: optional leakage model was not run",
            "body": [
                "The guide proposes a raw-versus-deduplicated cross-validation experiment.",
                "The project scope excludes model fitting; no AUC from that experiment is reported.",
                "Current leakage evidence comes from row order, duplicates, and ID-only AUC.",
            ],
            "notes": "This optional experiment is intentionally not run under the governing no-model-training constraint. Do not quote an unmeasured raw or cleaned model score.",
        },
    ]
    return slides, backup


def add_slide(prs, content, index, total, hidden=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if hidden:
        slide._element.set("show", "0")
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(249, 251, 253)

    title_box = slide.shapes.add_textbox(Inches(0.65), Inches(0.35), Inches(12.0), Inches(0.85))
    title_frame = title_box.text_frame
    title_frame.clear()
    title_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    title_p = title_frame.paragraphs[0]
    title_p.text = content["title"]
    title_p.font.name = "Aptos Display"
    title_p.font.size = Pt(29)
    title_p.font.bold = True
    title_p.font.color.rgb = RGBColor(24, 51, 74)

    body_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.45), Inches(11.8), Inches(4.9))
    body_frame = body_box.text_frame
    body_frame.clear()
    body_frame.word_wrap = True
    body_frame.margin_left = Inches(0.12)
    body_frame.margin_right = Inches(0.12)
    body_frame.margin_top = Inches(0.12)
    body_frame.margin_bottom = Inches(0.12)

    figures = [path for path in content_figures(content) if path.exists()]
    if figures:
        body_box.height = Inches(1.0)
        body_frame.paragraphs[0].text = content["body"][0]
        body_frame.paragraphs[0].font.size = Pt(20)
        body_frame.paragraphs[0].font.color.rgb = RGBColor(55, 70, 82)
        column_width = 11.1 / len(figures)
        for figure_index, figure_path in enumerate(figures):
            with Image.open(figure_path) as image:
                image_width, image_height = image.size
            scale = min(column_width / image_width, 3.4 / image_height)
            display_width = image_width * scale
            display_height = image_height * scale
            left = 1.1 + figure_index * column_width + (column_width - display_width) / 2
            slide.shapes.add_picture(
                str(figure_path),
                Inches(left),
                Inches(2.3),
                width=Inches(display_width),
                height=Inches(display_height),
            )
        remaining = content["body"][1:]
        if remaining:
            note_box = slide.shapes.add_textbox(Inches(0.85), Inches(5.9), Inches(11.5), Inches(0.85))
            note_paragraph = note_box.text_frame.paragraphs[0]
            note_paragraph.text = "  |  ".join(remaining)
            note_paragraph.font.size = Pt(20)
            note_paragraph.font.color.rgb = RGBColor(55, 70, 82)
    else:
        body_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        for item_index, item in enumerate(content["body"]):
            paragraph = body_frame.paragraphs[0] if item_index == 0 else body_frame.add_paragraph()
            paragraph.text = item
            paragraph.level = 0
            paragraph.font.name = "Aptos"
            paragraph.font.size = Pt(21 if len(content["body"]) <= 4 else 20)
            paragraph.font.color.rgb = RGBColor(45, 61, 74)
            paragraph.space_after = Pt(18)

    footer = slide.shapes.add_textbox(Inches(0.65), Inches(7.05), Inches(12.0), Inches(0.25))
    footer_frame = footer.text_frame
    footer_frame.clear()
    footer_p = footer_frame.paragraphs[0]
    footer_p.text = f"Research prototype - not a diagnostic tool                                      {index} / {total}"
    footer_p.font.name = "Aptos"
    footer_p.font.size = Pt(10)
    footer_p.font.color.rgb = RGBColor(100, 112, 124)
    footer_p.alignment = PP_ALIGN.RIGHT

    notes = slide.notes_slide.notes_text_frame
    notes.text = content["notes"]
    return slide


def build_pdf(slides, output, total):
    with PdfPages(output) as pdf:
        for index, content in enumerate(slides, 1):
            fig = plt.figure(figsize=(13.333, 7.5), facecolor="#f9fbfd")
            ax = fig.add_axes([0.06, 0.08, 0.88, 0.84])
            ax.axis("off")
            ax.text(
                0,
                0.98,
                content["title"],
                transform=ax.transAxes,
                va="top",
                fontsize=27,
                fontweight="bold",
                color="#18334a",
                wrap=True,
            )
            body = "\n".join(f"- {line}" for line in content["body"])
            figures = [path for path in content_figures(content) if path.exists()]
            if figures:
                ax.text(
                    0.01,
                    0.79,
                    body,
                    transform=ax.transAxes,
                    va="top",
                    fontsize=16,
                    color="#374652",
                    linespacing=1.25,
                )
                column_width = 0.88 / len(figures)
                for figure_index, figure_path in enumerate(figures):
                    image = plt.imread(figure_path)
                    ratio = image.shape[1] / image.shape[0]
                    width = min(column_width, 0.48 * ratio / 0.84)
                    height = width * 0.84 / ratio
                    left = 0.06 + figure_index * column_width + (column_width - width) / 2
                    image_ax = fig.add_axes([left, 0.18, width, height])
                    image_ax.imshow(image)
                    image_ax.axis("off")
            else:
                ax.text(
                    0.03,
                    0.78,
                    body,
                    transform=ax.transAxes,
                    va="top",
                    fontsize=21,
                    color="#2d3d4a",
                    linespacing=1.7,
                )
            ax.text(
                1,
                -0.03,
                f"Research prototype - not a diagnostic tool     {index} / {total}",
                transform=ax.transAxes,
                ha="right",
                va="bottom",
                fontsize=10,
                color="#64707c",
            )
            pdf.savefig(fig)
            plt.close(fig)


def build():
    if not RESULTS.exists():
        raise FileNotFoundError(f"Run `python -m src.results` first: {RESULTS}")
    data = json.loads(RESULTS.read_text(encoding="utf-8"))
    SLIDES_DIR.mkdir(parents=True, exist_ok=True)
    main, backup = build_deck(data)
    total = len(main) + len(backup)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    markdown = ["# Presentation source and speaker notes", ""]

    for index, content in enumerate(main, 1):
        add_slide(prs, content, index, total)
        markdown.extend([
            f"## Slide {index}: {content['title']}",
            "",
            *[f"- {line}" for line in content["body"]],
            "",
            *(
                [f"**Figure references:** {', '.join(content['figure_refs'])}", ""]
                if content.get("figure_refs")
                else []
            ),
            f"**Speaker notes:** {content['notes']}",
            "",
        ])
    for offset, content in enumerate(backup, len(main) + 1):
        add_slide(prs, content, offset, total, hidden=True)
        markdown.extend([
            f"## Hidden backup {offset - len(main)}: {content['title']}",
            "",
            *[f"- {line}" for line in content["body"]],
            "",
            f"**Speaker notes:** {content['notes']}",
            "",
        ])

    (SLIDES_DIR / "slides.md").write_text("\n".join(markdown), encoding="utf-8")
    output = SLIDES_DIR / "presentation.pptx"
    prs.save(output)
    pdf_path = SLIDES_DIR / "presentation.pdf"
    build_pdf(main, pdf_path, total)
    if len(main) != data["deliverable_spec"]["main_slides"]:
        raise AssertionError(f"Expected {data['deliverable_spec']['main_slides']} main slides, got {len(main)}")
    if len(backup) != data["deliverable_spec"]["backup_slides"]:
        raise AssertionError(f"Expected {data['deliverable_spec']['backup_slides']} backup slides, got {len(backup)}")
    hidden_count = sum(slide._element.get("show") == "0" for slide in prs.slides)
    if len(prs.slides) != total or hidden_count != len(backup):
        raise AssertionError("The PPTX main/hidden slide structure does not match the build.")
    if any(not slide.notes_slide.notes_text_frame.text.strip() for slide in prs.slides):
        raise AssertionError("Every slide must include speaker notes.")
    body_sizes = [
        paragraph.font.size.pt
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame and "Research prototype - not a diagnostic tool" not in shape.text
        for paragraph in shape.text_frame.paragraphs
        if paragraph.font.size
    ]
    if not body_sizes or min(body_sizes) < 20:
        raise AssertionError("Slide body text must be at least 20 pt.")
    pdf_pages = len(re.findall(rb"/Type\s*/Page\b", pdf_path.read_bytes()))
    if pdf_pages != len(main):
        raise AssertionError(f"Slide PDF has {pdf_pages} pages; expected {len(main)}.")
    print(f"Built {output}: {len(main)} main slides, {len(backup)} hidden backup slides")
    print(f"Slide PDF pages: {pdf_pages}; minimum body text: {min(body_sizes):.0f} pt")


if __name__ == "__main__":
    build()
